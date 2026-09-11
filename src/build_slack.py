#!/usr/bin/env python3
"""Build the AgentLab human-in-the-loop slide.

NOTE: AGENTLAB-SLACK-HANDOFF.md is out of date on this. It describes a bridge
that appends every channel message to the campaign boards, with the secretary
alongside. The current design, per the user:

  - the secretary sees messages first
  - it puts one on ANNOUNCEMENTS.md only when an agent has to answer
  - a campaign agent both posts its status and answers questions itself

Who answers from what:
  secretary       the workspace files
  campaign agent  its live context

slack_thread.png is a real thread from the user's workspace.

Output: agentlab_slack.pptx
"""

import os

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from slidekit import (
    ROOT, IMAGES,
    round_picture, new_deck, blank_slide, title_block, box, label_box, text,
    arrow, MONO, MUTED, FAINT, BLUE, BLUE_BG, VIOLET, VIOLET_BG, TEAL, TEAL_BG,
    SLATE, GREY_BG, HAIRLINE, ARROW,
)

SHOT = os.path.join(IMAGES, "slack_thread.png")
SHOT_ASPECT = 720 / 283

def add(prs):
    slide = blank_slide(prs)

    title_block(
        slide,
        "Human in the loop",
        "The lab runs on its own. You can still watch it, question it, "
        "and redirect it — from Slack.",
    )

    COL_X, COL_W = 0.50, 2.30
    SLACK_X, SLACK_W = 4.08, 1.21
    # the screenshot sets the width of the right column: it is shown at its
    # native 720 px / 99.7 px-per-inch, so its type matches the slide's.
    RIGHT_X, RIGHT_W = 5.61, 7.22
    MID = COL_X + COL_W / 2
    GAP_L, GAP_R = COL_X + COL_W, SLACK_X
    ROWS = [2.28, 3.63, 4.98]

    # ------------------------------------------------------------ the channel
    sl = box(slide, SLACK_X, 1.80, SLACK_W, 3.60, fill=VIOLET_BG, stroke=VIOLET,
             line_w=1.75)
    label_box(sl, [("Slack", 15, VIOLET, True, False),
                   ("team members", 9.5, MUTED, False, False),
                   ("and agents", 9.5, MUTED, False, False)])

    # ------------------------------------------------------------- the pieces
    sec = box(slide, COL_X, ROWS[0] - 0.43, COL_W, 0.86, fill=TEAL_BG,
              stroke=TEAL, line_w=1.75)
    label_box(sec, [("Secretary", 12.5, TEAL, True, False),
                    ("answers from workspace files", 8.5, MUTED, False, False)])

    board = box(slide, COL_X, ROWS[1] - 0.43, COL_W, 0.86, fill=GREY_BG,
                stroke=SLATE, line_w=1.25)
    label_box(board, [("ANNOUNCEMENTS.md", 10, SLATE, True, False),
                      ("the campaign's board", 8.5, MUTED, False, False)])
    board.text_frame.paragraphs[0].font.name = MONO

    agent = box(slide, COL_X, ROWS[2] - 0.43, COL_W, 0.86, fill=BLUE_BG,
                stroke=BLUE, line_w=1.75)
    label_box(agent, [("Campaign agent", 12.5, BLUE, True, False),
                      ("answers from its live context", 8.5, MUTED, False, False)])

    # --------------------------------------------------------------- traffic
    # the secretary reads the channel and replies into it
    arrow(slide, GAP_L, ROWS[0], GAP_R, ROWS[0], head=True, tail=True)
    text(slide, GAP_L - 0.02, ROWS[0] - 0.36, GAP_R - GAP_L + 0.04, 0.26,
         [("sees every message", 8, ARROW, False, False)])

    # and hands one on only when an agent has to answer it
    arrow(slide, MID, ROWS[0] + 0.43, MID, ROWS[1] - 0.43, color=SLATE, width=1.5)
    text(slide, MID + 0.12, ROWS[0] + 0.46, 1.60, 0.36,
         [("only when an agent", 8, MUTED, False, False),
          ("has to answer", 8, MUTED, False, False)], align=PP_ALIGN.LEFT)

    arrow(slide, MID, ROWS[1] + 0.43, MID, ROWS[2] - 0.43, color=SLATE, width=1.5)
    text(slide, MID + 0.12, ROWS[1] + 0.52, 1.60, 0.26,
         [("read between rounds", 8, MUTED, False, False)], align=PP_ALIGN.LEFT)

    # the agent posts its own status and its own answers
    arrow(slide, GAP_L, ROWS[2], GAP_R, ROWS[2])
    text(slide, GAP_L - 0.02, ROWS[2] - 0.36, GAP_R - GAP_L + 0.04, 0.26,
         [("Status / Answers", 8.5, ARROW, False, False)])

    # a standalone note, not part of the Slack path at all
    note = box(slide, RIGHT_X, 1.60, RIGHT_W, 0.54, fill=GREY_BG, stroke=HAIRLINE,
               line_w=1.0, radius=0.14)
    label_box(note, [("You can put an announcement in ANNOUNCEMENTS.md for the "
                      "agents at any time.", 10, SLATE, False, False)])

    # ------------------------------------------------------------ the thread
    SHOT_X, SHOT_W = RIGHT_X, RIGHT_W
    SHOT_H = SHOT_W / SHOT_ASPECT
    SHOT_Y = 2.40
    pic = slide.shapes.add_picture(SHOT, Inches(SHOT_X), Inches(SHOT_Y),
                                   Inches(SHOT_W), Inches(SHOT_H))
    pic.line.color.rgb = HAIRLINE
    pic.line.width = Pt(1.0)
    round_picture(pic)

    text(slide, SHOT_X, SHOT_Y + SHOT_H + 0.14, SHOT_W, 0.28,
         [("an alert, a start notice, a question, and the secretary answering it",
           9.5, FAINT, False, True)], align=PP_ALIGN.LEFT)

    # ------------------------------------------------------------- footnote
    text(slide, 0.5, 6.20, 12.33, 0.32,
         [("The secretary answers from the workspace files; the research agents "
           "answer from their live context.", 10, MUTED, False, False)],
         align=PP_ALIGN.LEFT)
    return slide


if __name__ == "__main__":
    deck = new_deck()
    add(deck)
    dest = os.path.join(ROOT, "agentlab_slack.pptx")
    deck.save(dest)
    print("wrote", dest)
