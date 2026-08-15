"""Shared drawing helpers and palette for the AgentLab slides.

White background, native editable shapes, 16:9.
"""

import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

FONT = "Calibri"
MONO = "Consolas"

# ---------------------------------------------------------------- palette
INK        = RGBColor(0x1A, 0x1A, 0x1A)
MUTED      = RGBColor(0x6B, 0x72, 0x80)
FAINT      = RGBColor(0x98, 0xA1, 0xB0)

BLUE       = RGBColor(0x1A, 0x56, 0xDB)   # the agent - the protagonist
BLUE_BG    = RGBColor(0xEF, 0xF4, 0xFF)
VIOLET     = RGBColor(0x6D, 0x28, 0xD9)   # LLM service
VIOLET_BG  = RGBColor(0xF6, 0xF3, 0xFF)
TEAL       = RGBColor(0x0F, 0x76, 0x6E)   # Globus Compute
TEAL_BG    = RGBColor(0xEC, 0xFD, 0xF7)
AMBER      = RGBColor(0xB4, 0x53, 0x09)   # compute systems
AMBER_BG   = RGBColor(0xFF, 0xF9, 0xED)
SLATE      = RGBColor(0x33, 0x41, 0x55)
GREY       = RGBColor(0x6B, 0x72, 0x80)
GREY_BG    = RGBColor(0xF7, 0xF7, 0xF8)
PANEL_BG   = RGBColor(0xFC, 0xFC, 0xFD)
HAIRLINE   = RGBColor(0xD5, 0xD9, 0xE0)
ARROW      = RGBColor(0x47, 0x55, 0x69)
RETURN     = RGBColor(0x05, 0x96, 0x69)   # the result path
BAR        = RGBColor(0x94, 0xA3, 0xB8)   # chart bars, unremarkable
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    return slide


def title_block(slide, title, subtitle):
    text(slide, 0.5, 0.30, 12.33, 0.55,
         [(title, 30, INK, True, False)], align=PP_ALIGN.LEFT)
    text(slide, 0.5, 0.86, 12.33, 0.34,
         [(subtitle, 13.5, MUTED, False, False)], align=PP_ALIGN.LEFT)


def _set_para(p, text, size, color, bold=False, italic=False,
              align=PP_ALIGN.CENTER, font=None):
    p.text = text
    p.alignment = align
    f = p.font
    f.name = font or FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def box(slide, x, y, w, h, *, fill, stroke, line_w=1.25, radius=0.06,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = stroke
    s.line.width = Pt(line_w)
    if dash:
        s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    s.text_frame.word_wrap = True
    s.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for attr, val in (("margin_left", 0.06), ("margin_right", 0.06),
                      ("margin_top", 0.03), ("margin_bottom", 0.03)):
        setattr(s.text_frame, attr, Inches(val))
    return s


def label_box(shape, lines):
    """lines = [(text, size, color, bold, italic), ...]"""
    tf = shape.text_frame
    for i, (t, size, color, bold, italic) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_para(p, t, size, color, bold, italic)
        if i > 0:
            p.space_before = Pt(2)


def text(slide, x, y, w, h, lines, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.TOP, font=None, line_space=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    for i, (t, size, color, bold, italic) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_para(p, t, size, color, bold, italic, align, font)
        if line_space:
            p.line_spacing = line_space
    return tb


def _add_ends(shape, head=True, tail=False, round_cap=False, size='med'):
    """headEnd/tailEnd are the last children of CT_LineProperties, so the
    library's own line formatting must be applied before these are appended."""
    ln = shape.line._get_or_add_ln()
    if round_cap:
        ln.set('cap', 'rnd')
    ends = {'type': 'triangle', 'w': size, 'len': size}
    if tail:
        ln.append(ln.makeelement(qn('a:headEnd'), dict(ends)))
    if head:
        ln.append(ln.makeelement(qn('a:tailEnd'), dict(ends)))


def arc_arrow(slide, cx, cy, r, a0, a1, *, color, width=5.0, segments=32,
              head=True, head_size='lg'):
    """A thick arc from a0 to a1, degrees clockwise from the top of the circle.

    Drawn as a freeform polyline rather than the CIRCULAR_ARROW preset, whose
    adjustment handles do not give control over the sweep.
    """
    pts = []
    for i in range(segments + 1):
        a = math.radians(a0 + (a1 - a0) * i / segments)
        pts.append((Inches(cx + r * math.sin(a)), Inches(cy - r * math.cos(a))))
    builder = slide.shapes.build_freeform(pts[0][0], pts[0][1])
    builder.add_line_segments(pts[1:], close=False)
    shape = builder.convert_to_shape()
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.shadow.inherit = False
    _add_ends(shape, head=head, round_cap=True, size=head_size)
    return shape


def round_picture(pic, adj=2600):
    """Round a picture's corners.

    python-pptx has no API for a picture's geometry, so swap the default
    prstGeom="rect" for "roundRect". In CT_ShapeProperties the geometry follows
    a:xfrm, so it is inserted there rather than appended.
    """
    spPr = pic._element.spPr
    for geom in spPr.findall(qn('a:prstGeom')):
        spPr.remove(geom)
    geom = spPr.makeelement(qn('a:prstGeom'), {'prst': 'roundRect'})
    avLst = geom.makeelement(qn('a:avLst'), {})
    avLst.append(avLst.makeelement(qn('a:gd'),
                                   {'name': 'adj', 'fmla': f'val {adj}'}))
    geom.append(avLst)
    xfrm = spPr.find(qn('a:xfrm'))
    spPr.insert(list(spPr).index(xfrm) + 1 if xfrm is not None else 0, geom)
    return pic


def arrow(slide, x1, y1, x2, y2, *, color=ARROW, width=1.75,
          head=True, tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color          # creates the <a:ln> element
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    # headEnd/tailEnd are the last children of CT_LineProperties, so append is safe
    if tail:
        ln.append(ln.makeelement(qn('a:headEnd'),
                                 {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    if head:
        ln.append(ln.makeelement(qn('a:tailEnd'),
                                 {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    return c
